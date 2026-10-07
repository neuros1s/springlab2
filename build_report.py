from pathlib import Path
import argparse, json, subprocess, textwrap
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'src/main/java/kz/iitu/springlab'
EVIDENCE=ROOT/'evidence'
OUT=ROOT/'submission'

def build(draft=False):
    OUT.mkdir(exist_ok=True)
    data={} if draft else json.loads((EVIDENCE/'results.json').read_text())
    if not draft and data.get('checksPassed') is not True:
        raise RuntimeError('The application checks have not passed. Refusing to produce a final report.')
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.65)
    sec.left_margin=sec.right_margin=Inches(.7)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Liberation Sans';st.font.color.rgb=RGBColor(0,0,0)
        for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
    st=doc.styles['Normal'];st.font.size=Pt(10.5);st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.line_spacing=1.05
    def p(t):doc.add_paragraph(t)
    def h(t):doc.add_heading(t,1)
    def new(t):doc.add_page_break();h(t)
    def code(t):
        q=doc.add_paragraph();q.paragraph_format.line_spacing=1
        q.paragraph_format.space_after=Pt(5)
        r=q.add_run(t.rstrip());r.font.name='Liberation Mono';r.font.size=Pt(8)
    def listing(file):
        doc.add_heading(Path(file).name.replace('.java',''),2)
        code((SRC/file).read_text())
    def table(head,rows):
        t=doc.add_table(rows=1,cols=len(head));t.autofit=True
        for c,v in zip(t.rows[0].cells,head):c.text=str(v)
        for row in rows:
            for c,v in zip(t.add_row().cells,row):c.text=str(v)
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
        t._tbl.tblPr.append(borders)
        for c in t.rows[0].cells:
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1');c._tc.get_or_add_tcPr().append(shade)
            for r in c.paragraphs[0].runs:r.bold=True
    def picture(name,title,expected):
        if draft:
            p('Expected observation: '+expected)
            p('Runtime evidence is pending. Run mvn clean verify and the report script to insert the actual output here.')
            return
        text=(EVIDENCE/(name+'.log')).read_text().strip()
        if name=='proxy':text=json.dumps(json.loads((EVIDENCE/'proxy.response.txt').read_text()),indent=2)
        if name=='success':
            n=0;lines=[]
            for line in text.splitlines():
                if any(tag in line for tag in ['[AUDIT]','[LOG]','[TIME]']):n+=1;line=f'{n}. '+line
                lines.append(line)
            text='\n'.join(lines)
        # This image displays captured output, not fabricated terminal chrome.
        wrapped=[]
        for line in text.splitlines():wrapped.extend(textwrap.wrap(line,95,replace_whitespace=False,drop_whitespace=False) or [''])
        fontpath='/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf'
        font=ImageFont.truetype(fontpath,20)
        titlefont=ImageFont.truetype(fontpath,22)
        image=Image.new('RGB',(1220,95+29*len(wrapped)), '#111827');draw=ImageDraw.Draw(image)
        draw.text((24,18),title,font=titlefont,fill='#ffffff')
        draw.text((24,50),'Actual Spring Boot integration test output',font=font,fill='#91a4bf')
        for i,line in enumerate(wrapped):draw.text((24,85+i*29),line,font=font,fill='#e6edf3')
        fp=EVIDENCE/(name+'.png');image.save(fp)
        width=min(6.8,6.6*image.width/image.height)
        doc.add_picture(str(fp),width=Inches(width))
        p(title+'. Captured from the tested application; text is laid out for readability.')
    doc.add_heading('Laboratory Work 4',0)
    h('Cross cutting concerns with Spring AOP')
    p('Course: RWPSF 3305 - Web Application Development with Spring Framework')
    p('Student: Turgynbek Amirbek\nGroup: IS2413\nClass-list position: 20\nIndividual assignment: Variant 8 - call tracing')
    p('The variants repeat after 12, so position 20 corresponds to variant 8: ((20 - 1) mod 12) + 1 = 8.')
    h('Aim')
    p('To move method logging, execution timing and user-action auditing into aspects, compare advice order, inspect the Spring proxy, and demonstrate why self-invocation bypasses interception. Variant 8 adds nested call tracing using Before and After advice.')
    p('Repository: https://github.com/neuros1s/springlab2')
    p('Target branch: https://github.com/neuros1s/springlab2/tree/lab04')
    if draft:p('Execution status: the code and checks are prepared; runtime results have not yet been captured. The branch link is the intended upload location.')
    else:
        p('Evidence source: '+data['source']+'. Requests are executed by MockMvc inside a real Spring application context, without opening a network port.')
        p('Capture time: '+data['capturedAt']+'\nJava version: '+data['javaVersion'])
    h('Build and project')
    p('The Maven project includes the Lab 2 and Lab 3 classes and adds the Lab 4 service, controller, annotation and aspects. Spring Boot 3.5.16 manages the AOP starter dependency. The compiler targets Java 25, as required by the assignment.')
    code('mvn clean verify\njava -jar target/spring-lab-01-0.0.1-SNAPSHOT.jar')
    new('Pointcuts and audit annotation')
    listing('aspect/Pointcuts.java');listing('audit/Audited.java')
    p('serviceLayer selects methods declared in the service package and its subpackages. publicMethod limits interception to public methods. serviceOperation combines both conditions. The aspects reuse this named pointcut. The audit aspect has one annotation-binding expression so it can receive the Audited instance.')
    p('RUNTIME retention makes the annotation available to Spring while the application is running. action identifies the operation; logArguments controls whether the argument values are recorded.')
    new('Logging advice');listing('aspect/LoggingAspect.java')
    p('Before records the method and arguments. AfterReturning records a successful result. AfterThrowing records the exception type and message. Its throwing name matches the ex parameter, and returning matches result. A successful call uses Before and AfterReturning; a failing call uses Before and AfterThrowing.')
    new('Execution timing');listing('aspect/TimingAspect.java')
    p('System.nanoTime measures elapsed time. proceed executes the next interceptor or target method and preserves the return value. The finally block records time on both success and failure. The deliberately slow findAll method sleeps for 300 ms, so it should exceed the 200 ms warning threshold.')
    new('Auditing and advice order');listing('aspect/AuditAspect.java')
    p('AuditAspect has order 1, LoggingAspect order 2, and TimingAspect order 3. Among these three, audit enters first and exits last. Exceptions are logged and rethrown. The controller maps invalid identifiers to HTTP 400 only after the exception has passed through the aspects.')
    table(['Position','Core record'],[(1,'AUDIT start'),(2,'LOG arguments'),(3,'TIME elapsed'),(4,'LOG result'),(5,'AUDIT success')])
    new('Individual assignment call tracing');listing('aspect/TracingAspect.java')
    p('TracingAspect has order 0, so its entry and exit surround the three core aspects. Each thread has its own depth value. Before prints the current indentation and increments the depth. After runs on success and failure, reduces the depth, and removes the ThreadLocal value when the outermost call finishes. This avoids stale indentation on reused request threads.')
    p('The two spaces per level make proxy-visible nesting readable. Self-invoked methods are intentionally absent from this trace because they never pass through a proxy.')
    new('Service and self invocation fix');listing('service/CatalogService.java');listing('service/FixedRemovalService.java')
    p('CatalogService.removeTwice uses this.remove and retains the original problem. FixedRemovalService is a separate coordinating bean; its two calls reach the injected CatalogService proxy. This separates orchestration from removal and avoids a circular self reference.')
    new('Successful and failing requests')
    p('GET /api/lab4/items?limit=5 returns five items. The numbered core records show audit, logging and timing in the required relative order; the tracing records wrap them.')
    picture('success','Successful catalog list','Five core records in order, plus tracing entry and exit.')
    p('DELETE /api/lab4/item/0 raises IllegalArgumentException. The aspects observe the failure and preserve it; the controller returns HTTP 400 with the type and message.')
    picture('failure','Invalid identifier','LOG exception, AUDIT failure, timing and tracing cleanup; no successful-return advice.')
    new('Proxy inspection and observations')
    picture('proxy','GET /api/lab4/proxy','CGLIB subclass of CatalogService; both AOP and CGLIB flags are true.')
    p('The runtime bean is a generated subclass, typically containing $$SpringCGLIB$$ in its name. Its superclass is CatalogService. Spring Boot uses class-based proxies by default; the suffix is generated infrastructure, not a class written by the student.')
    def observed(name,token):
        return 'Pending' if draft else sum(token in line for line in (EVIDENCE/(name+'.log')).read_text().splitlines())
    table(['Call or measurement','Expected','Observed'],[
        ('External remove audit records',2,observed('external','[AUDIT]')),
        ('Self-invocation audit records',0,observed('self-before','[AUDIT]')),
        ('Self-invocation logging records',2,observed('self-before','[LOG]')),
        ('Fixed nested-call audit records',4,observed('self-after','[AUDIT]')),
        ('Fixed logging records',6,observed('self-after','[LOG]'))])
    p('A pair of audit records is start plus outcome for one intercepted operation. The two logging records before the fix belong to removeTwice itself, not to the two remove calls. After the fix the three intercepted operations create three logging pairs.')
    new('Self invocation before the fix')
    code('GET /api/lab4/remove-twice/5')
    picture('self-before','Direct internal calls','Only CatalogService.removeTwice is logged, timed and traced. There are no audit records for its internal remove calls.')
    p('A Spring proxy intercepts calls arriving through the proxy. Once execution enters the target, this refers to the target object; a neighbouring method call is a normal Java call. An annotation alone cannot make that internal call cross the proxy boundary.')
    new('Self invocation after the fix')
    code('GET /api/lab4/remove-twice-fixed/5')
    picture('self-after','Calls through another bean','The coordinator has depth 0 and the two CatalogService.remove calls have depth 1. Each removal is audited.')
    p('The integration check also calls the fixed endpoint with identifier 0, then makes a fresh request. The error must unwind both trace levels and the next request must start at depth 0. This checks cleanup rather than only the normal path.')
    new('Defence answers')
    answers=[
      ('1 What is a cross cutting concern','A responsibility needed across multiple operations. In this project it is logging, timing, auditing and tracing.'),
      ('2 Pointcut and advice','The Pointcuts methods answer where interception applies. The advice methods answer what action is executed at those join points.'),
      ('3 Meaning of the package expression','within(kz.iitu.springlab.service..*) matches types in service and nested packages. The double dot includes subpackages; a star is a wildcard within the relevant name pattern.'),
      ('4 Why two annotations','Aspect defines the aspect metadata. Component lets component scanning register an instance as a Spring bean.'),
      ('5 Advice order within one aspect','For these advice types the entry side is Around then Before. A normal or exceptional outcome runs AfterReturning or AfterThrowing; After has finally semantics. Control then returns to the outer Around advice.'),
      ('6 Removing proceed','The target is not called by that advice. The result is whatever Around returns; an exception occurs only if the advice itself throws.'),
      ('7 Missing AfterReturning on failure','The method threw instead of returning normally, so AfterThrowing applies.'),
      ('8 The throwing attribute','It binds the thrown exception to the advice parameter. A wrong parameter name prevents correct binding and normally causes aspect creation to fail.')]
    for q,a in answers:doc.add_heading(q,2);p(a)
    new('Defence answers continued')
    answers=[
      ('9 Runtime annotation retention','Spring must inspect the annotation while the application runs.'),
      ('10 Annotation binding','The annotation pointcut binds the matched Audited instance to the advice argument named audited.'),
      ('11 Which aspect runs first','Tracing has order 0 and wraps everything. Among the three required aspects, audit has order 1 and enters first and closes last.'),
      ('12 Generated class name','It identifies the CGLIB subclass used to intercept calls while delegating to the target implementation.'),
      ('13 Why CGLIB','Spring Boot defaults to class-based proxies. Merely adding an interface does not override that Boot setting.'),
      ('14 Self invocation and the fix','Internal calls bypass the proxy. A separate coordinator bean now calls the injected CatalogService proxy.'),
      ('15 Private and static methods','Private methods cannot be overridden by a subclass proxy. Static methods belong to the class rather than a proxied instance. Spring proxy-based AOP cannot advise either.'),
      ('16 Other annotations using proxies','Transactional, Async and Cacheable commonly use proxy-based interception when their corresponding support is enabled.')]
    for q,a in answers:doc.add_heading(q,2);p(a)
    h('Conclusions')
    p('Aspects keep repeated operational concerns out of catalog methods. Named pointcuts make the interception boundary explicit, while Order makes nesting predictable. Proxy-based interception explains why internal calls behave differently from calls between beans. Separating the coordinator from the removal bean restores auditing and exposes nested calls to tracing.')
    p('References: Laboratory work No. 4 assignment; Spring Framework Reference, Declaring Advice and Proxying Mechanisms.')
    file=OUT/('Lab04_Turgynbek_Amirbek_IS2413_DRAFT.docx' if draft else 'Lab04_Turgynbek_Amirbek_IS2413.docx')
    doc.save(file);print(file)

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--draft',action='store_true');build(parser.parse_args().draft)
