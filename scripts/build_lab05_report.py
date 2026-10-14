from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / 'src/main/java/kz/iitu/springlab'
EVIDENCE = ROOT / 'evidence/lab05'
OUT = ROOT / 'submission'
DATA = json.loads((EVIDENCE / 'results.json').read_text(encoding='utf-8'))
if DATA.get('checksPassed') is not True:
    raise RuntimeError('A final report requires successful live HTTP and browser checks')
if any(DATA['mavenTests'][k] for k in ('failures','errors','skipped')):
    raise RuntimeError('Maven checks did not pass')
REQ = DATA['requests']

def header(row, name):
    return next((v for k,v in row['headers'].items() if k.lower()==name.lower()), '')

doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.65)
for name in ['Normal','Title','Heading 1','Heading 2']:
    style=doc.styles[name]
    style.font.name='Liberation Sans'
    style.font.color.rgb=RGBColor(0,0,0)
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
normal=doc.styles['Normal']
normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(28)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)
doc.core_properties.author='Turgynbek Amirbek'
doc.core_properties.title='Laboratory Work 5 Spring MVC'
doc.core_properties.subject='REST controllers views Post Redirect Get and language headers'

def p(text):
    doc.add_paragraph(text)
def h(text):
    doc.add_heading(text,1)
def new(text):
    doc.add_page_break()
    h(text)
def code(text, size=8.1):
    q=doc.add_paragraph()
    q.paragraph_format.space_after=Pt(7)
    q.paragraph_format.line_spacing=1
    r=q.add_run(text.rstrip())
    r.font.name='Liberation Mono'
    r.font.size=Pt(size)
def listing(relative):
    doc.add_heading(Path(relative).stem,2)
    code((SRC/relative).read_text(encoding='utf-8'))
def table(head, rows, widths=None):
    t=doc.add_table(rows=1,cols=len(head))
    t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,text in zip(t.rows[0].cells,head):
        c.text=str(text)
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6')
        c._tc.get_or_add_tcPr().append(shade)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,text in zip(t.add_row().cells,row): c.text=str(text)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single')
        b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths: c.width=Inches(widths[j])
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(4)
                para.paragraph_format.space_before=Pt(4)
                para.paragraph_format.line_spacing=1.05
                for run in para.runs:
                    run.font.size=Pt(9)
                    if i==0:run.bold=True
    doc.add_paragraph()
def picture(file, caption):
    fp=EVIDENCE/file
    with Image.open(fp) as img: width=min(6.95,6.4*img.width/img.height)
    doc.add_picture(str(fp),width=Inches(width))
    q=doc.add_paragraph(caption)
    q.paragraph_format.space_after=Pt(8)
    for r in q.runs:r.font.size=Pt(9)
def readable_body(name):
    raw=REQ[name]['body']
    if not raw:return 'Empty body'
    try:
        data=json.loads(raw)
        if isinstance(data,list):
            return '\n'.join(json.dumps(x,ensure_ascii=False) for x in data) or '[]'
        return json.dumps(data,ensure_ascii=False)
    except ValueError:
        return raw

doc.add_heading('Laboratory Work 5',0)
h('Developing a web application with Spring MVC')
p('Controllers routing request parameters models views and static resources')
p('JSC International Information Technology University\n'
  'Faculty of Business Media and Management\nDepartment of Information Systems')
p('Course: RWPSF 3305 - Web Application Development with Spring Framework')
p('Student: Turgynbek Amirbek\nGroup: IS2413\nClass-list position: 20\n'
  'Individual assignment: Variant 8 - Accept-Language and JSON field names')
p('Position 20 corresponds to variant 8 when the 12 variants repeat.')
p('Branch: https://github.com/neuros1s/springlab2/tree/lab05')
p('Evidence capture: '+DATA['capturedAt'])

new('Aim and project setup')
p('The application exposes a book catalogue as a JSON API and as an HTML page. '
  'The work covers URL mapping, request argument binding, ResponseEntity status codes, '
  'the model, Thymeleaf view resolution, static files and Post Redirect Get. '
  'The individual variant changes the field names of the list response through the '
  'Accept-Language header.')
p('The project continues the tested Lab 4 implementation. The lab05 branch starts from '
  'lab04 while its pull request awaits review; the existing Lab 2, Lab 3 and Lab 4 code '
  'remains available in the same application.')
table(['Component','Used in this run'],[
    ('JDK',DATA['javaVersion'].splitlines()[0]),
    ('Spring Boot','3.5.16'),
    ('Dependencies','Web, Thymeleaf and the existing validation, AOP and test starters'),
    ('Automated Java tests',str(DATA['mavenTests']['tests'])+' passed, no failures, errors or skips'),
    ('Runtime verification',DATA['source'])
],[1.7,5.45])
code('mvn clean verify\njava -jar target/spring-lab-01-0.0.1-SNAPSHOT.jar\n'
     'http://localhost:8080/books')
p('Each run begins with three records. The repository keeps data in memory, so restarting '
  'the application restores the seed collection. No database credentials are required.')
h('How requests are processed')
p('DispatcherServlet receives the request. HandlerMapping selects a method by its URL, '
  'HTTP method and media conditions. HandlerAdapter resolves and converts the arguments '
  'before invoking the controller. The REST result goes to a JSON message converter; '
  'the page result goes to the view resolver and Thymeleaf.')
table(['Annotation','Source in the request','Example'],[
    ('PathVariable','URL path segment','id in /api/books/1'),
    ('RequestParam','Query string or form field','author and limit; title, author and year in the form'),
    ('RequestHeader','HTTP header','Accept-Language'),
    ('RequestBody','The complete request body','JSON Book submitted by POST or PUT')
],[1.55,2.55,3.05])

new('Domain and repository')
listing('catalog/Book.java')
listing('catalog/BookRepository.java')
p('Book is an immutable record. ConcurrentHashMap holds records and AtomicLong assigns '
  'identifiers. The list is ordered by identifier so the demonstration is reproducible. '
  'Replacement uses computeIfPresent and cannot insert a missing identifier.')

new('Service layer')
listing('catalog/BookService.java')
p('The author filter matches a case-insensitive fragment, so Bloch finds Joshua Bloch. '
  'A creation always receives a new identifier, even if the client supplies one. PUT '
  'uses the path identifier and changes all data fields. Titles and authors must be '
  'nonblank and the year must be between 1 and 9999.')

rest=(SRC/'web/BookRestController.java').read_text(encoding='utf-8')
cut=rest.index('    @GetMapping("/{id}")')
new('REST controller and language header')
code(rest[:cut])
p('The list method reads two query parameters and one header. The default limit is 10. '
  'A negative limit is rejected with 400; a nonnumeric limit fails during argument binding. '
  'The controller declares a JSON representation, so XML is not available.')
new('REST controller continued')
code(rest[cut:])
p('ResponseEntity supplies the status, headers and body. POST uses 201 and a Location URI '
  'for the assigned identifier. PUT returns 200 only for an existing record. DELETE has '
  'no response body and returns 204 once; a second deletion returns 404.')

new('MVC controller and the model')
listing('web/BookPageController.java')
p('The GET handler adds books, author and title to the model. Returning books/list is a '
  'logical view name, not text sent to the browser. The POST handler reads form parameters, '
  'creates one record and redirects to a GET. Its flash attribute is available to that GET '
  'and then removed.')
new('Template and static resources')
code((ROOT/'src/main/resources/templates/books/list.html').read_text(encoding='utf-8'),7.8)
p('The template is stored under templates/books/list.html. The stylesheet is '
  'static/css/app.css and the welcome page is static/index.html. Spring Boot supplies '
  'the resource handler and welcome-page mapping; neither static URL has a custom controller.')

new('Observed REST operations')
p('The HTTP observations below were captured from the running JAR. The page form was '
  'tested first, so the initial API list also contains Java Precisely.')
table(['Operation and request','Status','Response body'],[
    ('List\nGET /api/books',REQ['list_all']['status'],readable_body('list_all')),
    ('Read\nGET /api/books/1',REQ['read_one']['status'],readable_body('read_one')),
    ('Create\nPOST /api/books',REQ['create']['status'],readable_body('create')),
    ('Replace\nPUT /api/books/1',REQ['replace']['status'],readable_body('replace')),
    ('Delete\nDELETE /api/books/2',REQ['delete_first']['status'],readable_body('delete_first'))
],[2,0.6,4.55])
p('Creation Location: '+header(REQ['create'],'Location'))
p('The POST body was {"title":"Pro Spring 6","author":"Cosmina","year":2023}. '
  'PUT supplied identifier 999 in the body while targeting /api/books/1; the result kept '
  'identifier 1, showing that the path is authoritative.')
table(['Additional check','Observed status','Result'],[
    ('GET /api/books?author=Bloch&limit=1',REQ['list_filtered']['status'],
     readable_body('list_filtered')),
    ('GET /api/books/999',REQ['read_absent']['status'],'No record found'),
    ('PUT /api/books/999',REQ['replace_absent']['status'],'No record inserted'),
    ('Second DELETE /api/books/2',REQ['delete_second']['status'],'The first DELETE already removed it')
],[3,0.75,3.4])
p('The first deletion returned an empty body. GET of the newly created Location returned '
  '200. GET of the absent PUT target still returned 404.')

new('Catalogue page')
picture('catalogue.png','Figure 1. GET /books before mutations, with the three seed records and the stylesheet applied.')
p('The live browser check also verified the computed border-collapse property of the '
  'table. GET /css/app.css returned '+str(REQ['stylesheet']['status'])+
  ' with a text/css content type, and the root welcome page returned '+str(REQ['home']['status'])+'.')
new('Author filter')
picture('filtered.png','Figure 2. GET /books?author=Bloch. Only the matching record is rendered.')
p('The filter changes the books collection in the model before the template iterates over '
  'it. The matching author is Joshua Bloch; the other two seed records are absent.')

new('Post Redirect Get and flash message')
picture('flash.png','Figure 3. The redirected GET after adding Java Precisely. The flash message appears once.')
prg=DATA['prg']
table(['Request or observation','Measured result'],[
    ('POST /books',str(prg['postStatus'])+'; Location '+prg['postLocation']),
    ('Redirected GET /books',prg['redirectedGetStatus']),
    ('Books before and after POST',str(prg['booksBefore'])+' -> '+str(prg['booksAfterPost'])),
    ('GET caused by refresh',str(prg['refreshStatus'])+'; '+str(prg['booksAfterRefresh'])+' books'),
    ('Flash message removed on refresh',prg['flashRemovedOnRefresh'])
],[3.45,3.7])
p('Refreshing repeats the GET, not the POST, so no second record is created. '
  'SessionFlashMapManager stores the flash attribute between the requests and removes '
  'it when the redirect target receives it. Ordinary model attributes are request-scoped.')
new('Catalogue after refreshing')
picture('refreshed.png','Figure 4. A second GET keeps four books and no longer displays the flash message.')

new('Status code experiment')
table(['Request','Observed status','Part that failed'],[
    ('GET /api/books/abc',REQ['error_400']['status'],
     'The path matched, but abc cannot be converted to the long id argument. Binding failed before the method ran.'),
    ('GET /api/nothing',REQ['error_404']['status'],
     'No API mapping exists for this path and no matching static resource exists.'),
    ('DELETE /api/books',REQ['error_405']['status'],
     'The collection path exists for GET and POST; DELETE requires the /{id} path. The HTTP method did not match.'),
    ('POST /api/books\nContent-Type: text/plain',REQ['error_415']['status'],
     'The POST mapping consumes application/json. The incoming body media type did not match.'),
    ('GET /api/books\nAccept: application/xml',REQ['error_406']['status'],
     'The API produces application/json and cannot supply the requested XML representation.')
],[2.5,0.65,4])
p('The two extra limit checks also returned 400: limit=abc failed conversion, while '
  'limit=-1 was rejected by the controller. These are different failure points even '
  'though they produce the same status.')
p('415 concerns the representation supplied by the client. 406 concerns the representation '
  'the client is willing to receive. The live response files are included with the report evidence.')

new('Individual variant 8')
p('GET /api/books reads Accept-Language through RequestHeader. English is the default. '
  'Russian changes title, author and year to название, автор and год; id remains the shared '
  'identifier. Regional tags such as ru-RU and weighted preferences are accepted. '
  'An unsupported language falls back to English.')
start=rest.index('    @GetMapping\n')
end=rest.index('    @GetMapping("/{id}")')
code(rest[start:end],8)
for name,label in [('variant_en','English request'),('variant_ru','Russian request')]:
    row=REQ[name]
    doc.add_heading(label,2)
    code(row['method']+' '+row['path']+'\nAccept-Language: '+
         row['requestHeaders']['Accept-Language']+'\nHTTP '+str(row['status'])+
         '\nContent-Language: '+header(row,'Content-Language')+
         '\nVary: '+header(row,'Vary')+'\n\n'+readable_body(name),8)
p('Both responses describe the same stored book. Content-Language identifies the selected '
  'language, and Vary: Accept-Language tells caches that this header affects the answer.')

new('Demonstration answers')
answers=[
('1 Trace the request','The browser sends an HTTP request to the embedded server. DispatcherServlet asks HandlerMapping for a handler, and HandlerAdapter resolves arguments and invokes it. ResponseEntity and the JSON converter produce an API answer; a view name and model produce an HTML answer.'),
('2 HandlerMapping and HandlerAdapter','HandlerMapping selects the matching handler. HandlerAdapter knows how to invoke that handler, including argument resolution and return-value processing.'),
('3 ResponseBody','ResponseBody makes the return value the response body instead of a logical view name. RestController includes this annotation. The page controller also uses it on its small validation error handler.'),
('4 Sources of list parameters','author and limit come from the query string through RequestParam. Accept-Language comes from a request header through RequestHeader. The same RequestParam annotation reads the fields of the browser form.'),
('5 A nonnumeric limit','Spring cannot convert the string to int and returns 400 before calling the list method. A negative numeric limit binds successfully but is rejected inside the method.'),
('6 Optional RequestParam','A required parameter missing from the request gives 400. required=false allows author to be absent; the service then returns the full collection. defaultValue supplies 10 when limit is absent or empty.'),
('7 Creation and Location','201 means that a new resource was created. Location contains /api/books/{assignedId}, which can be requested with GET. A client-supplied id does not overwrite an existing record.'),
('8 Deleting twice','The first call removes the record and returns 204 with no body. The next call cannot find that record and returns 404. Repeating the deletion leaves the same absent state.'),
('9 The model','The model is the data passed to a view for the current request. The page controller fills books, author and title; Thymeleaf reads them to render the heading, filter and table.'),
('10 Resolving the view name','The Thymeleaf view resolver uses its template prefix and suffix to resolve books/list to classpath:/templates/books/list.html. The controller depends on the logical name rather than the template file syntax.'),
('11 Serving CSS','Spring Boot configures a static resource handler for classpath:/static. /css/app.css resolves to static/css/app.css. static/index.html is the welcome page for the root URL.'),
('12 Requests in Post Redirect Get','The form produces POST /books with 302. The browser follows Location with GET /books and receives 200. Refresh adds another GET; it does not repeat the creation. CSS requests are separate resource requests.'),
('13 Flash storage','RedirectAttributes adds the message to the output FlashMap. SessionFlashMapManager temporarily stores it, matches the target request, and makes it available to that request model. It then removes the saved flash map.'),
('14 Mapping failures','400 in the required experiment is an argument conversion failure. 404 is an unknown path or missing resource. 405 is the wrong method for an existing path. 415 rejects the incoming content type; 406 rejects an unavailable outgoing representation.'),
('15 415 and 406','text/plain does not meet the consumes application/json condition of POST. Accept application/xml does not meet the JSON produces condition of GET. The failure happens before normal controller processing.'),
('16 Replacing the template engine','The REST controller does not need to change. The page controller can keep its model and logical view name if the new view resolver uses the same name. The template syntax and resolver configuration would change.')
]
for i,(q,a) in enumerate(answers):
    if i in [6,12]:new('Demonstration answers continued')
    doc.add_heading(q,2);p(a)
h('Conclusions')
p('The front controller gives the application one consistent entry point for web requests. '
  'Handler mapping and argument binding keep routing and conversion outside the domain code. '
  'View resolution and message converters let the same service support both HTML and JSON. '
  'Post Redirect Get prevents duplicate form submissions, while explicit response statuses '
  'make successful operations and mapping failures clear.')
p('References: Laboratory Work No. 5 assignment; Spring Framework 6.2 Reference, '
  'Mapping Requests and Flash Attributes; Spring Boot Reference, Servlet Web Applications.')
p('https://docs.spring.io/spring-framework/reference/6.2/web/webmvc/mvc-controller/ann-requestmapping.html\n'
  'https://docs.spring.io/spring-framework/reference/6.2/web/webmvc/mvc-controller/ann-methods/flash-attributes.html')
OUT.mkdir(exist_ok=True)
target=OUT/'Lab05_Turgynbek_Amirbek_IS2413.docx'
doc.save(target)
print(target)
